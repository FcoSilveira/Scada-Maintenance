# -*- coding: utf-8 -*-

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

BASE_DIR = Path(__file__).resolve().parent
ARQUIVO_CSV = BASE_DIR / "entrada" / "alarmes.csv"
ARQUIVO_RELATORIO = BASE_DIR / "saida" / "relatorio_alarmes.xlsx"

ENCODINGS_CSV = ("utf-8-sig", "utf-8", "cp1252", "cp850", "iso-8859-1")

LIMITE_OK = 5
LIMITE_ATENCAO = 15

MAPA_NIVEIS = {
    "1": "INFORMATION",
    "2": "IMPORTANT",
    "3": "URGENT",
    "4": "CRITICAL",
}


def definir_status(quantidade):
    if quantidade <= LIMITE_OK:
        return "OK"
    if quantidade <= LIMITE_ATENCAO:
        return "ATENCAO"
    return "CRITICO"


def carregar_dados():
    print("Lendo CSV...")

    ultimo_erro = None
    for encoding in ENCODINGS_CSV:
        try:
            df = pd.read_csv(ARQUIVO_CSV, sep=";", encoding=encoding)
            df.attrs["encoding"] = encoding
            break
        except UnicodeDecodeError as erro:
            ultimo_erro = erro
    else:
        raise Exception(f"Nao foi possivel ler o CSV. Ultimo erro: {ultimo_erro}")

    df.columns = df.columns.str.strip()

    colunas_obrigatorias = ["DataHora", "Tipo", "xid", "pointName", "alarmLevel", "Mensagem"]
    colunas_faltantes = [col for col in colunas_obrigatorias if col not in df.columns]
    if colunas_faltantes:
        raise Exception(f"CSV invalido. Colunas faltantes: {colunas_faltantes}")

    if df.empty:
        raise Exception("O arquivo CSV de alarmes esta vazio.")

    print(f"CSV lido com encoding: {df.attrs['encoding']}")
    return df


def processar_alarmes(df_bruto):
    print("Processando alarmes...")

    df = df_bruto[["Tipo", "xid", "pointName", "alarmLevel", "Mensagem"]].copy()
    df.columns = ["Tipo", "XID", "PointName", "AlarmLevel", "Mensagem"]

    df["Tipo"] = df["Tipo"].fillna("").astype(str).str.strip()
    df["XID"] = df["XID"].fillna("").astype(str).str.strip()
    df["PointName"] = df["PointName"].fillna("").astype(str).str.strip()
    df["Mensagem"] = df["Mensagem"].fillna("").astype(str).str.strip()
    df["AlarmLevel"] = (
        df["AlarmLevel"].fillna("").astype(str).str.strip().str.upper().replace(MAPA_NIVEIS)
    )

    df_agrupado = (
        df.groupby(["Tipo", "XID", "PointName", "Mensagem", "AlarmLevel"], dropna=False)
        .size()
        .reset_index(name="Quantidade")
    )

    df_agrupado["Status"] = df_agrupado["Quantidade"].apply(definir_status)
    return df_agrupado.sort_values(by="Quantidade", ascending=False).reset_index(drop=True)


def montar_resumo(df_todos, total_registros):
    print("Gerando estatisticas...")

    xids_preenchidos = df_todos.loc[df_todos["XID"] != "", "XID"]
    total_xids = xids_preenchidos.nunique()
    total_mensagens = df_todos["Mensagem"].nunique()
    maior_ocorrencia = int(df_todos["Quantidade"].max())
    media_ocorrencias = round(float(df_todos["Quantidade"].mean()), 2)

    agg_xid = (
        df_todos[df_todos["XID"] != ""]
        .groupby("XID", as_index=False)
        .agg(Quantidade_Total=("Quantidade", "sum"))
        .sort_values(by="Quantidade_Total", ascending=False)
    )
    xid_mais_alarmes = agg_xid.iloc[0]["XID"] if not agg_xid.empty else "N/A"

    agg_mensagem = (
        df_todos.groupby("Mensagem", as_index=False)
        .agg(Quantidade_Total=("Quantidade", "sum"))
        .sort_values(by="Quantidade_Total", ascending=False)
    )
    mensagem_mais_recorrente = agg_mensagem.iloc[0]["Mensagem"] if not agg_mensagem.empty else "N/A"

    return pd.DataFrame(
        {
            "Metrica": [
                "Total de registros",
                "Total de XIDs diferentes",
                "Total de mensagens diferentes",
                "Maior quantidade de ocorrencias",
                "Quantidade media de ocorrencias",
                "XID com maior quantidade de alarmes",
                "Mensagem mais recorrente",
            ],
            "Valor": [
                total_registros,
                total_xids,
                total_mensagens,
                maior_ocorrencia,
                media_ocorrencias,
                xid_mais_alarmes,
                mensagem_mais_recorrente,
            ],
        }
    )


def montar_top_alarmes(df_todos):
    return (
        df_todos.groupby("Mensagem", as_index=False)
        .agg(
            Quantidade_Total=("Quantidade", "sum"),
            Quantidade_de_XIDs=("XID", lambda valores: valores[valores != ""].nunique()),
        )
        .sort_values(by="Quantidade_Total", ascending=False)
        .rename(columns={"Quantidade_de_XIDs": "Quantidade de XIDs"})
    )


def montar_top_xids(df_todos):
    df_com_xid = df_todos[df_todos["XID"] != ""].copy()
    if df_com_xid.empty:
        return pd.DataFrame(columns=["XID", "PointName", "Quantidade Total", "Mensagens Diferentes"])

    return (
        df_com_xid.groupby(["XID", "PointName"], as_index=False)
        .agg(Quantidade_Total=("Quantidade", "sum"), Mensagens_Diferentes=("Mensagem", "nunique"))
        .sort_values(by="Quantidade_Total", ascending=False)
        .rename(columns={"Quantidade_Total": "Quantidade Total", "Mensagens_Diferentes": "Mensagens Diferentes"})
    )


def montar_por_nivel(df_todos):
    return (
        df_todos.groupby("AlarmLevel", as_index=False)
        .agg(Quantidade_Total=("Quantidade", "sum"), Mensagens_Diferentes=("Mensagem", "nunique"))
        .sort_values(by="Quantidade_Total", ascending=False)
        .rename(columns={"Quantidade_Total": "Quantidade Total", "Mensagens_Diferentes": "Mensagens Diferentes"})
    )


def montar_por_tipo(df_todos):
    return (
        df_todos.groupby("Tipo", as_index=False)
        .agg(Quantidade_Total=("Quantidade", "sum"), Mensagens_Diferentes=("Mensagem", "nunique"))
        .sort_values(by="Quantidade_Total", ascending=False)
        .rename(columns={"Quantidade_Total": "Quantidade Total", "Mensagens_Diferentes": "Mensagens Diferentes"})
    )


def aplicar_formatacao(arquivo):
    print("Aplicando formatacao...")

    try:
        wb = load_workbook(arquivo)
    except PermissionError as erro:
        raise PermissionError("Feche o arquivo de relatorio no Excel e execute novamente.") from erro

    verde = PatternFill("solid", fgColor="C6EFCE")
    amarelo = PatternFill("solid", fgColor="FFEB9C")
    vermelho = PatternFill("solid", fgColor="FFC7CE")

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)

        if ws.title == "Resumo":
            ws.column_dimensions["A"].width = 40
            ws.column_dimensions["B"].width = 45
        elif ws.title == "Todos":
            ws.column_dimensions["A"].width = 15
            ws.column_dimensions["B"].width = 30
            ws.column_dimensions["C"].width = 40
            ws.column_dimensions["D"].width = 45
            ws.column_dimensions["E"].width = 18
            ws.column_dimensions["F"].width = 15
            ws.column_dimensions["G"].width = 15

            for row in range(2, ws.max_row + 1):
                status = ws.cell(row, 7).value
                if status == "OK":
                    fill = verde
                elif status == "ATENCAO":
                    fill = amarelo
                elif status == "CRITICO":
                    fill = vermelho
                else:
                    continue

                for col in range(1, ws.max_column + 1):
                    ws.cell(row, col).fill = fill
        elif ws.title == "Top Alarmes":
            ws.column_dimensions["A"].width = 45
            ws.column_dimensions["B"].width = 20
            ws.column_dimensions["C"].width = 22
        elif ws.title == "Top XIDs":
            ws.column_dimensions["A"].width = 30
            ws.column_dimensions["B"].width = 40
            ws.column_dimensions["C"].width = 18
            ws.column_dimensions["D"].width = 22
        elif ws.title in ("Por Nivel", "Por Tipo"):
            ws.column_dimensions["A"].width = 22
            ws.column_dimensions["B"].width = 20
            ws.column_dimensions["C"].width = 22

    try:
        wb.save(arquivo)
    except PermissionError as erro:
        raise PermissionError("Feche o arquivo de relatorio no Excel e execute novamente.") from erro


def exportar_relatorio(df_resumo, df_todos, df_top_alarmes, df_top_xids, df_por_nivel, df_por_tipo):
    print("Gerando relatorio Excel...")

    ARQUIVO_RELATORIO.parent.mkdir(parents=True, exist_ok=True)
    temp_path = ARQUIVO_RELATORIO.with_name("relatorio_alarmes_temp.xlsx")

    if temp_path.exists():
        temp_path.unlink()

    try:
        with pd.ExcelWriter(temp_path, engine="openpyxl") as writer:
            df_resumo.to_excel(writer, sheet_name="Resumo", index=False)
            df_todos.to_excel(writer, sheet_name="Todos", index=False)
            df_top_alarmes.to_excel(writer, sheet_name="Top Alarmes", index=False)
            df_top_xids.to_excel(writer, sheet_name="Top XIDs", index=False)
            df_por_nivel.to_excel(writer, sheet_name="Por Nivel", index=False)
            df_por_tipo.to_excel(writer, sheet_name="Por Tipo", index=False)

        if ARQUIVO_RELATORIO.exists():
            ARQUIVO_RELATORIO.unlink()
        temp_path.replace(ARQUIVO_RELATORIO)
    except PermissionError as erro:
        raise PermissionError("Feche o arquivo de relatorio no Excel e execute novamente.") from erro


def main():
    df_bruto = carregar_dados()
    df_todos = processar_alarmes(df_bruto)

    total_registros = len(df_bruto)

    df_resumo = montar_resumo(df_todos, total_registros)
    df_top_alarmes = montar_top_alarmes(df_todos)
    df_top_xids = montar_top_xids(df_todos)
    df_por_nivel = montar_por_nivel(df_todos)
    df_por_tipo = montar_por_tipo(df_todos)

    exportar_relatorio(df_resumo, df_todos, df_top_alarmes, df_top_xids, df_por_nivel, df_por_tipo)
    aplicar_formatacao(ARQUIVO_RELATORIO)

    print("\n" + "=" * 50)
    print("        AUDITORIA DE ALARMES OWEN CLOUD")
    print("=" * 50)
    print(f"Registros analisados..........: {total_registros}")
    print(f"XIDs diferentes...............: {df_resumo.iloc[1, 1]}")
    print(f"Mensagens diferentes..........: {df_resumo.iloc[2, 1]}")
    print(f"Maior ocorrencia..............: {df_resumo.iloc[3, 1]}")
    print(f"XID com mais alarmes..........: {df_resumo.iloc[5, 1]}")
    print(f"Mensagem mais recorrente......: {df_resumo.iloc[6, 1]}")
    print("=" * 50)
    print(f"Relatorio.....................: {ARQUIVO_RELATORIO}")
    print("=" * 50)
    print("Finalizado.")


if __name__ == "__main__":
    main()
